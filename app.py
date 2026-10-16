import io
import re
from urllib.parse import urlparse
import streamlit as st
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="AEM Translation URL Builder", page_icon="🌐", layout="wide")

st.markdown("""
<style>
.block-container {padding-top: 2rem; padding-bottom: 3rem; max-width: 1500px;}
.hero {padding: 2rem 2.2rem; border-radius: 24px; background: linear-gradient(135deg,#071a2c,#0b5d5b); color:white; margin-bottom:1.2rem; box-shadow:0 12px 35px rgba(0,0,0,.14)}
.hero h1 {font-size:2.35rem; margin:0 0 .35rem 0;}
.hero p {opacity:.88; margin:0; font-size:1.05rem}
.pill {display:inline-block; padding:.32rem .68rem; margin:.18rem .18rem .18rem 0; border-radius:999px; background:rgba(255,255,255,.12); border:1px solid rgba(255,255,255,.2); font-size:.88rem}
[data-testid="stMetric"] {background:rgba(127,127,127,.07); border:1px solid rgba(127,127,127,.15); padding:14px; border-radius:18px;}
</style>
<div class="hero"><h1>🌐 AEM Translation URL Builder</h1><p>Upload a Thermo Fisher translation request workbook. The app reads only the first sheet, detects selected languages (X), and generates Local Author, Regional Author, Regional Sites/Rollout, and Global Author URLs.</p><br><span class="pill">First sheet only</span><span class="pill">Automatic locale detection</span><span class="pill">Wrapped Excel URLs</span><span class="pill">One Excel sheet per language</span><span class="pill">Flags by language</span></div>
""", unsafe_allow_html=True)

AUTHOR="https://author-prod-use1.aemprod.thermofisher.net"
LOCALES={
"de-DE": {"flag":"🇩🇪","label":"German / Germany","root":"europe","local":"de-de","regional":"en-de"},
"es-ES": {"flag":"🇪🇸","label":"Spanish / Spain","root":"europe","local":"es-es","regional":"en-es"},
"fr-FR": {"flag":"🇫🇷","label":"French / France","root":"europe","local":"fr-fr","regional":"en-fr"},
"ja-JP": {"flag":"🇯🇵","label":"Japanese / Japan","root":"japan","local":"ja-jp","regional":"en-jp"},
"ko-KR": {"flag":"🇰🇷","label":"Korean / Korea","root":"ipac","local":"ko-kr","regional":"en-kr"},
"zh-CN": {"flag":"🇨🇳","label":"Chinese / Mainland China","root":"greater-china","local":"zh-cn","regional":"en-cn"},
"pt-BR": {"flag":"🇧🇷","label":"Portuguese / Brazil","root":"latin-america","local":"pt-br","regional":"en-br"},
"zh-TW": {"flag":"🇹🇼","label":"Traditional Chinese / Taiwan","root":"ipac","local":"zh-tw","regional":"en-tw"},
"es-LATAM": {"flag":"🌎","label":"Spanish / Latin America","special":True},
}
LATAM={
"Argentina":{"flag":"🇦🇷","local":"es-ar","regional":"en-ar"},
"Chile":{"flag":"🇨🇱","local":"es-cl","regional":"en-cl"},
"Mexico":{"flag":"🇲🇽","local":"es-mx","regional":"en-mx"},
}

def clean_header(v):
    return str(v or '').split('\n')[0].strip()

def page_path(url):
    p=urlparse(str(url).strip()).path
    m=re.search(r'/home/(.+)$',p)
    if not m: return None
    return 'home/'+m.group(1)

def urls_for(path, cfg):
    root,local,regional=cfg['root'],cfg['local'],cfg['regional']
    nohtml=path[:-5] if path.endswith('.html') else path
    return {
        "Local Author":f"{AUTHOR}/editor.html/content/lifetech/{root}/{local}/{path}",
        "Regional Author":f"{AUTHOR}/editor.html/content/lifetech/{root}/{regional}/{path}",
        "Regional Sites / Rollout":f"{AUTHOR}/sites.html/content/lifetech/{root}/{regional}/{nohtml}",
        "Global Author":f"{AUTHOR}/editor.html/content/lifetech/global/en/{path}",
    }

def parse(file_bytes):
    wb=load_workbook(io.BytesIO(file_bytes),data_only=False)
    ws=wb[wb.sheetnames[0]]
    header_row=None; url_col=None
    for r in range(1,min(ws.max_row,20)+1):
        for c in range(1,ws.max_column+1):
            if 'URL in AEM' in str(ws.cell(r,c).value or ''):
                header_row,url_col=r,c; break
        if header_row: break
    if not header_row: raise ValueError("Could not find the 'URL in AEM' header in the first sheet.")
    lang_cols={}
    for c in range(1,ws.max_column+1):
        h=clean_header(ws.cell(header_row,c).value)
        if h in LOCALES: lang_cols[h]=c
    rows=[]
    for r in range(header_row+1,ws.max_row+1):
        url=ws.cell(r,url_col).value
        if not url or not str(url).strip().startswith(('http://','https://')): continue
        title=ws.cell(r,max(1,url_col-1)).value or ''
        selected=[lang for lang,c in lang_cols.items() if str(ws.cell(r,c).value or '').strip().lower()=='x']
        if selected: rows.append((r,str(title),str(url).strip(),selected))
    return ws.title,rows

def style_url_sheet(ws):
    navy="0B1F33"; white="FFFFFF"; green="008000"
    headers=["Flag","Language","Page Title","Original URL","Local Author","Regional Author","Regional Sites / Rollout","Global Author"]
    ws.append(headers)
    ws.sheet_view.showGridLines=False
    for cell in ws[1]:
        cell.fill=PatternFill('solid',fgColor=navy)
        cell.font=Font(color=white,bold=True)
        cell.alignment=Alignment(vertical='center',wrap_text=True)
    ws.row_dimensions[1].height=34
    widths=[8,28,34,58,70,70,70,70]
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    ws.freeze_panes="A2"


def append_url_row(ws, flag, language, title, original, generated):
    ws.append([flag,language,title,original,generated['Local Author'],generated['Regional Author'],generated['Regional Sites / Rollout'],generated['Global Author']])
    r=ws.max_row
    thin=Side(style='thin',color='D8E2E8')
    for c in ws[r]:
        c.alignment=Alignment(vertical='top',wrap_text=True)
        c.border=Border(bottom=thin)
    ws.cell(r,1).alignment=Alignment(horizontal='center',vertical='top',wrap_text=True)
    for col in range(4,9):
        c=ws.cell(r,col)
        c.font=Font(color='008000',underline='single')
        c.hyperlink=c.value


def make_excel(rows, latam_countries):
    wb=Workbook()
    wb.remove(wb.active)

    detected=[]
    for _,_,_,langs in rows:
        for lang in langs:
            if lang not in detected: detected.append(lang)

    # One worksheet per selected language. LATAM stays together in one sheet,
    # with a row for each selected destination country.
    for lang in detected:
        sheet_name=lang[:31]
        ws=wb.create_sheet(sheet_name)
        style_url_sheet(ws)
        for _,title,original,langs in rows:
            if lang not in langs: continue
            path=page_path(original)
            if not path: continue
            if lang=='es-LATAM':
                for country in latam_countries:
                    x=LATAM[country]
                    cfg={"root":"latin-america","local":x['local'],"regional":x['regional']}
                    u=urls_for(path,cfg)
                    append_url_row(ws,x['flag'],f"es-LATAM — {country}",title,original,u)
            else:
                cfg=LOCALES[lang]
                u=urls_for(path,cfg)
                append_url_row(ws,cfg['flag'],f"{lang} — {cfg['label']}",title,original,u)
        ws.auto_filter.ref=ws.dimensions

    info=wb.create_sheet("Locale Map")
    info.sheet_view.showGridLines=False
    info.append(["Language","Flag","Local root","Regional / Rollout root"])
    for lang,cfg in LOCALES.items():
        if lang=='es-LATAM': continue
        info.append([lang,cfg['flag'],f"/content/lifetech/{cfg['root']}/{cfg['local']}",f"/content/lifetech/{cfg['root']}/{cfg['regional']}"])
    for country,x in LATAM.items():
        info.append([f"es-LATAM — {country}",x['flag'],f"/content/lifetech/latin-america/{x['local']}",f"/content/lifetech/latin-america/{x['regional']}"])
    for c in info[1]:
        c.fill=PatternFill('solid',fgColor='0B6B68')
        c.font=Font(color='FFFFFF',bold=True)
    for col,w in enumerate([24,10,52,52],1): info.column_dimensions[get_column_letter(col)].width=w
    for row in info.iter_rows():
        for c in row: c.alignment=Alignment(vertical='top',wrap_text=True)
    info.freeze_panes="A2"

    out=io.BytesIO(); wb.save(out); return out.getvalue()

uploaded=st.file_uploader("Drop your translation request Excel here",type=["xlsx","xlsm"],help="Only the first worksheet is analyzed.")
if uploaded:
    data=uploaded.getvalue()
    try: sheet,rows=parse(data)
    except Exception as e: st.error(str(e)); st.stop()
    langs=sorted({l for *_,ls in rows for l in ls})
    c1,c2,c3=st.columns(3); c1.metric("First sheet",sheet); c2.metric("Pages with selections",len(rows)); c3.metric("Languages detected",len(langs))
    st.markdown("### Detected languages")
    st.markdown(" ".join(f"**{LOCALES[l]['flag']} {l}**" for l in langs))
    latam_countries=list(LATAM)
    if 'es-LATAM' in langs:
        st.info("es-LATAM is selected. Choose the Latin America destinations to generate.")
        latam_countries=st.multiselect("LATAM destinations",list(LATAM),default=list(LATAM))
    with st.expander("Preview detected pages",expanded=True):
        for _,title,url,ls in rows:
            st.markdown(f"**{title or 'Untitled page'}**  \n{url}  \n"+" · ".join(f"{LOCALES[l]['flag']} {l}" for l in ls))
    output=make_excel(rows,latam_countries)
    st.success("Your URL workbook is ready — each detected language has its own worksheet.")
    st.download_button("⬇️ Download generated Excel",output,file_name="aem_translation_urls.xlsx",mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",use_container_width=True)
else:
    st.info("Upload an .xlsx file to begin. The original workbook is not modified.")
